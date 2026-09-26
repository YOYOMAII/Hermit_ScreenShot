#!/usr/bin/env python3
"""Insert code screenshots into a Word document, two figures per page."""

from __future__ import annotations

import re
import string
from dataclasses import dataclass, field
from pathlib import Path
from typing import IO, Dict, List, Optional, Sequence, Tuple, Union
from xml.sax.saxutils import escape, quoteattr

try:
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn
    from docx.shared import Inches
    from docx.text.paragraph import Paragraph
except ImportError as exc:
    raise SystemExit(
        "Missing dependency. Install with: python -m pip install -r requirements.txt"
    ) from exc


DEFAULT_CAPTION = "Fig {n}: {stem}{part}{ext}"
CAPTION_FIELDS = frozenset({"n", "file_no", "part", "stem", "ext", "file"})
INSERT_MODES = ("continue", "after_page", "end")
FIGURES_PER_PAGE = 2
CAPTION_FONT = "Arial"
CAPTION_HALF_POINTS = 22  # Word stores font sizes in half-points: 22 = 11 pt.
CAPTION_HEIGHT = Inches(0.4)
SLOT_GAP = Inches(0.15)
TAG_PREFIX = "HermitSmart Fig "
TAG_PATTERN = re.compile(r"HermitSmart Fig (\d+)$")

NAMESPACES = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}
DECLARATIONS = " ".join(f'xmlns:{prefix}="{uri}"' for prefix, uri in NAMESPACES.items())

Source = Union[str, Path, IO[bytes]]


@dataclass(frozen=True)
class Figure:
    """One screenshot with the details used to build its caption."""

    image: Path
    stem: str
    ext: str
    file_no: int
    part: int


@dataclass
class _Break:
    block: int
    hard: bool
    at_start: bool = False
    at_end: bool = False


@dataclass
class _Scan:
    blocks: List[object]
    breaks: List[_Break] = field(default_factory=list)
    figures: Dict[int, int] = field(default_factory=dict)  # figure number -> block index
    last_block: Optional[int] = None
    max_doc_pr_id: int = 0

    @property
    def pages_estimate(self) -> int:
        return len(self.breaks) + 1

    @property
    def next_figure(self) -> int:
        return max(self.figures, default=0) + 1


def caption_for(template: str, figure: Figure, number: int) -> str:
    """Fill the caption template, allowing only the documented placeholders."""
    allowed = ", ".join(f"{{{key}}}" for key in sorted(CAPTION_FIELDS))
    try:
        names = [name for _, name, _, _ in string.Formatter().parse(template)]
    except ValueError as exc:
        raise ValueError(f"Caption template is not valid ({exc}). "
                         "To show a brace, type it twice: {{ or }}.") from exc
    if any(name is not None and name not in CAPTION_FIELDS for name in names):
        raise ValueError(f"Caption template can only use {allowed}.")
    values = {
        "n": number, "file_no": figure.file_no, "part": figure.part,
        "stem": figure.stem, "ext": figure.ext, "file": f"{figure.stem}{figure.ext}",
    }
    try:
        return template.format(**values)
    except (ValueError, IndexError, KeyError, AttributeError) as exc:
        raise ValueError(f"Caption template is not valid. Use only {allowed}.") from exc


def _body_blocks(body) -> List[object]:
    return [child for child in body if child.tag != qn("w:sectPr")]


def _section_breaks_page(sect_pr) -> bool:
    kind = sect_pr.find(qn("w:type"))
    return kind is None or kind.get(qn("w:val"), "nextPage") != "continuous"


def _scan(document) -> _Scan:
    """Find page boundaries, HermitSmart figures, and the largest drawing ID.

    A .docx file has no page numbers, so pages are estimated from hard breaks
    and from the page-break markers Word saves after laying out the document.
    A rendered marker right after a hard break is the same page boundary.
    """
    body = document.element.body
    scan = _Scan(_body_blocks(body))
    content_since_break = False
    text_tags = {qn("w:t"), qn("w:drawing"), qn("w:pict"), qn("w:tab"), qn("w:object")}
    for index, block in enumerate(scan.blocks):
        # Document-order markers: a _Break, or None for visible content.
        markers: List[Optional[_Break]] = []
        section_break = False
        for element in block.iter():
            tag = element.tag
            if tag == qn("w:pageBreakBefore") and element.get(qn("w:val"), "1") not in ("0", "false"):
                markers.append(_Break(index, hard=True))
                content_since_break = False
            elif tag == qn("w:br") and element.get(qn("w:type")) == "page":
                markers.append(_Break(index, hard=True))
                content_since_break = False
            elif tag == qn("w:lastRenderedPageBreak"):
                if content_since_break:
                    markers.append(_Break(index, hard=False))
                content_since_break = False
            elif tag == qn("w:sectPr") and element.getparent().tag == qn("w:pPr"):
                section_break = _section_breaks_page(element)
            elif tag in text_tags and (tag != qn("w:t") or (element.text or "").strip()):
                markers.append(None)
                content_since_break = True
            elif tag == qn("wp:docPr"):
                try:
                    scan.max_doc_pr_id = max(scan.max_doc_pr_id, int(element.get("id", "0")))
                except ValueError:
                    pass
                match = TAG_PATTERN.match(element.get("name", ""))
                if match:
                    scan.figures[int(match.group(1))] = index
                    scan.last_block = index
        for position, event in enumerate(markers):
            if event is None:
                continue
            event.at_start = all(marker is not None for marker in markers[:position])
            event.at_end = all(marker is not None for marker in markers[position + 1:])
            scan.breaks.append(event)
        if section_break:
            scan.breaks.append(_Break(index, hard=True, at_end=True))
            content_since_break = False
    return scan


def inspect_document(source: Source) -> dict:
    """Summarise a document for the upload page."""
    scan = _scan(Document(source))
    return {
        "pages_estimate": scan.pages_estimate,
        "figures": len(scan.figures),
        "next_figure": scan.next_figure,
    }


def _starts_with_hard_break(scan: _Scan, index: int) -> bool:
    return any(event.block == index and event.hard and event.at_start for event in scan.breaks)


def _ends_with_hard_break(scan: _Scan, index: int) -> bool:
    return any(event.block == index and event.hard and event.at_end for event in scan.breaks)


def _has_content(blocks: Sequence[object]) -> bool:
    for block in blocks:
        if block.tag == qn("w:tbl"):
            return True
        for element in block.iter(qn("w:t"), qn("w:drawing"), qn("w:pict"), qn("w:object")):
            if element.tag != qn("w:t") or (element.text or "").strip():
                return True
    return False


def _geometry(section) -> dict:
    page_width = section.page_width or Inches(8.5)
    page_height = section.page_height or Inches(11)
    left = section.left_margin if section.left_margin is not None else Inches(1)
    right = section.right_margin if section.right_margin is not None else Inches(1)
    top = section.top_margin if section.top_margin is not None else Inches(1)
    bottom = section.bottom_margin if section.bottom_margin is not None else Inches(1)
    width = int(page_width - left - right)
    height = int(page_height - top - bottom)
    if width <= 0 or height <= CAPTION_HEIGHT * FIGURES_PER_PAGE:
        raise ValueError("The document's page margins leave no room for images.")
    slot = height // FIGURES_PER_PAGE
    return {"width": width, "slot": slot,
            "image_height": int(slot - CAPTION_HEIGHT - SLOT_GAP)}


def _anchor_attributes(z_order: int) -> str:
    # behindDoc="0" together with wrapNone is Word's "In Front of Text" layout.
    return (f'distT="0" distB="0" distL="0" distR="0" simplePos="0" '
            f'relativeHeight="{z_order}" behindDoc="0" locked="0" '
            f'layoutInCell="1" allowOverlap="1"')


def _position(x: int, y: int, cx: int, cy: int) -> str:
    return (f'<wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="margin"><wp:posOffset>{x}</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="margin"><wp:posOffset>{y}</wp:posOffset></wp:positionV>'
            f'<wp:extent cx="{cx}" cy="{cy}"/>'
            f'<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>')


def _image_run(rel_id: str, doc_pr_id: int, number: int, caption: str, filename: str,
               x: int, y: int, cx: int, cy: int, z_order: int):
    return parse_xml(
        f'<w:r {DECLARATIONS}><w:drawing>'
        f'<wp:anchor {_anchor_attributes(z_order)}>{_position(x, y, cx, cy)}'
        f'<wp:docPr id="{doc_pr_id}" name="{TAG_PREFIX}{number}" descr={quoteattr(caption)}/>'
        f'<wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>'
        f'<a:graphic><a:graphicData uri="{NAMESPACES["pic"]}"><pic:pic>'
        f'<pic:nvPicPr><pic:cNvPr id="0" name={quoteattr(filename)}/><pic:cNvPicPr/></pic:nvPicPr>'
        f'<pic:blipFill><a:blip r:embed="{rel_id}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
        f'</pic:pic></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>'
    )


def _caption_run(doc_pr_id: int, number: int, caption: str,
                 x: int, y: int, cx: int, cy: int, z_order: int):
    font = (f'<w:rPr><w:rFonts w:ascii="{CAPTION_FONT}" w:hAnsi="{CAPTION_FONT}" '
            f'w:eastAsia="{CAPTION_FONT}" w:cs="{CAPTION_FONT}"/>'
            f'<w:sz w:val="{CAPTION_HALF_POINTS}"/><w:szCs w:val="{CAPTION_HALF_POINTS}"/></w:rPr>')
    return parse_xml(
        f'<w:r {DECLARATIONS}><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing>'
        f'<wp:anchor {_anchor_attributes(z_order)}>{_position(x, y, cx, cy)}'
        f'<wp:docPr id="{doc_pr_id}" name="HermitSmart Caption {number}"/>'
        f'<wp:cNvGraphicFramePr/>'
        f'<a:graphic><a:graphicData uri="{NAMESPACES["wps"]}"><wps:wsp>'
        f'<wps:cNvSpPr txBox="1"/>'
        f'<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></wps:spPr>'
        f'<wps:txbx><w:txbxContent><w:p><w:pPr><w:spacing w:before="0" w:after="0"/>'
        f'<w:jc w:val="center"/>{font}</w:pPr>'
        f'<w:r>{font}<w:t xml:space="preserve">{escape(caption)}</w:t></w:r></w:p></w:txbxContent></wps:txbx>'
        f'<wps:bodyPr rot="0" vert="horz" wrap="square" lIns="91440" tIns="45720" '
        f'rIns="91440" bIns="45720" anchor="t" anchorCtr="0"><a:noAutofit/></wps:bodyPr>'
        f'</wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing>'
        f'</mc:Choice></mc:AlternateContent></w:r>'
    )


def _new_page_paragraph(document, break_before: bool):
    paragraph = Paragraph(parse_xml(f'<w:p {DECLARATIONS}/>'), document._body)
    paragraph.paragraph_format.page_break_before = break_before or None
    paragraph.paragraph_format.space_before = 0
    paragraph.paragraph_format.space_after = 0
    return paragraph._p


def _insertion_index(scan: _Scan, mode: str, after_page: int) -> int:
    if mode == "continue" and scan.last_block is not None:
        return scan.last_block + 1
    if mode == "after_page":
        if after_page <= 0:
            return 0
        if after_page < scan.pages_estimate:
            boundary = scan.breaks[after_page - 1]
            return boundary.block if boundary.at_start else boundary.block + 1
    return len(scan.blocks)


def build_document(
    source: Source,
    figures: Sequence[Figure],
    output: Union[str, Path],
    *,
    insert_mode: str = "end",
    after_page: int = 0,
    start_figure: Optional[int] = None,
    caption_template: str = DEFAULT_CAPTION,
) -> List[dict]:
    """Add figures to the document and save it; return each figure's caption."""
    if insert_mode not in INSERT_MODES:
        raise ValueError("Choose a valid insert position.")
    if not figures:
        raise ValueError("Add at least one code file.")
    document = Document(source)
    body = document.element.body
    scan = _scan(document)
    geometry = _geometry(document.sections[-1])
    number = start_figure if start_figure is not None else scan.next_figure
    if number < 1:
        raise ValueError("Start figure number must be 1 or more.")
    captions = [caption_for(caption_template, figure, number + offset)
                for offset, figure in enumerate(figures)]

    index = _insertion_index(scan, insert_mode, after_page)
    # Figures fill a half-empty HermitSmart page before new pages are made.
    slots: List[Tuple[object, int]] = []
    queue = list(range(len(figures)))
    if insert_mode == "continue" and scan.last_block is not None:
        last = scan.blocks[scan.last_block]
        used = sum(1 for block in scan.figures.values() if block == scan.last_block)
        if last.tag == qn("w:p") and used < FIGURES_PER_PAGE:
            for slot in range(used, FIGURES_PER_PAGE):
                if queue:
                    slots.append((last, slot))
                    queue.pop(0)

    before = scan.blocks[:index]
    first_break = bool(slots) or (_has_content(before) and not _ends_with_hard_break(scan, index - 1))
    new_pages = []
    while queue:
        page = _new_page_paragraph(document, break_before=first_break or bool(new_pages))
        new_pages.append(page)
        for slot in range(FIGURES_PER_PAGE):
            if queue:
                slots.append((page, slot))
                queue.pop(0)

    anchor = scan.blocks[index] if index < len(scan.blocks) else body.find(qn("w:sectPr"))
    for page in new_pages:
        if anchor is not None:
            anchor.addprevious(page)
        else:
            body.append(page)

    # Content that followed the insertion point starts on its own page again.
    if new_pages and index < len(scan.blocks) and not _starts_with_hard_break(scan, index):
        following = scan.blocks[index]
        if following.tag == qn("w:p"):
            Paragraph(following, document._body).paragraph_format.page_break_before = True
        else:
            following.addprevious(_new_page_paragraph(document, break_before=True))

    doc_pr_id = scan.max_doc_pr_id
    z_order = 251659264 + doc_pr_id
    result = []
    for offset, (figure, (paragraph, slot)) in enumerate(zip(figures, slots)):
        figure_number = number + offset
        caption = captions[offset]
        rel_id, image = document.part.get_or_add_image(str(figure.image))
        scale = min(geometry["width"] / image.px_width, geometry["image_height"] / image.px_height)
        cx, cy = int(image.px_width * scale), int(image.px_height * scale)
        x = (geometry["width"] - cx) // 2
        y = slot * geometry["slot"]
        doc_pr_id += 2
        z_order += 2
        paragraph.append(_image_run(rel_id, doc_pr_id - 1, figure_number, caption,
                                    figure.image.name, x, y, cx, cy, z_order - 1))
        paragraph.append(_caption_run(doc_pr_id, figure_number, caption,
                                      x, y + cy, cx, int(CAPTION_HEIGHT), z_order))
        result.append({"n": figure_number, "caption": caption, "page_slot": slot + 1})

    document.save(str(output))
    return result


__all__ = ["DEFAULT_CAPTION", "Figure", "INSERT_MODES", "build_document",
           "caption_for", "inspect_document"]
