"""Checks for inserting code screenshots into Word documents."""

import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from PIL import Image

import app as web_app
from hermit_smart import NAMESPACES, Figure, build_document, caption_for, inspect_document


def body_text(document):
    return [paragraph.text for paragraph in document.paragraphs if paragraph.text]


def page_paragraphs(document):
    return [paragraph for paragraph in document.paragraphs
            if paragraph._p.findall(".//" + qn("wp:anchor"))]


def docx_bytes(document):
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


class BuildDocumentTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.folder = Path(temp.name)
        self.images = []
        for index, height in enumerate((300, 200, 120), start=1):
            path = self.folder / f"index{index}.png"
            Image.new("RGB", (1200, height), "#eff0f1").save(path, dpi=(144, 144))
            self.images.append(path)
        self.figures = [Figure(path, "index", ".html", 1, part)
                        for part, path in enumerate(self.images, start=1)]

    def document_with(self, *texts, page_breaks=False):
        document = Document()
        for number, text in enumerate(texts):
            if page_breaks and number:
                document.add_page_break()
            document.add_paragraph(text)
        path = self.folder / "source.docx"
        document.save(path)
        return path

    def test_two_figures_per_page_in_front_of_text_with_arial_captions(self):
        output = self.folder / "out.docx"
        placed = build_document(self.document_with("Intro"), self.figures, output)
        self.assertEqual([item["caption"] for item in placed],
                         ["Fig 1: index1.html", "Fig 2: index2.html", "Fig 3: index3.html"])
        self.assertEqual([item["page_slot"] for item in placed], [1, 2, 1])

        document = Document(output)
        pages = page_paragraphs(document)
        self.assertEqual(len(pages), 2)
        self.assertTrue(all(page.paragraph_format.page_break_before for page in pages))
        anchors = document.element.body.findall(".//" + qn("wp:anchor"))
        self.assertEqual(len(anchors), 6)
        for anchor in anchors:
            self.assertEqual(anchor.get("behindDoc"), "0")
            self.assertIsNotNone(anchor.find(qn("wp:wrapNone")))
        boxes = document.element.body.findall(f".//{{{NAMESPACES['wps']}}}txbx")
        self.assertEqual(len(boxes), 3)
        captions = []
        for box in boxes:
            run = box.find(".//" + qn("w:r"))
            fonts = run.find(".//" + qn("w:rFonts"))
            self.assertEqual(fonts.get(qn("w:ascii")), "Arial")
            self.assertEqual(run.find(".//" + qn("w:sz")).get(qn("w:val")), "22")
            captions.append(run.find(qn("w:t")).text)
        self.assertEqual(captions, [item["caption"] for item in placed])
        self.assertEqual(inspect_document(output),
                         {"pages_estimate": 3, "figures": 3, "next_figure": 4})

    def test_continue_fills_the_half_empty_page_first(self):
        first = self.folder / "first.docx"
        build_document(self.document_with("Intro"), self.figures, first)
        second = self.folder / "second.docx"
        placed = build_document(first, self.figures[:2], second, insert_mode="continue")
        self.assertEqual([(item["n"], item["page_slot"]) for item in placed], [(4, 2), (5, 1)])
        document = Document(second)
        pages = page_paragraphs(document)
        self.assertEqual(len(pages), 3)
        tags = [len(page._p.findall(".//" + qn("wp:docPr"))) for page in pages]
        self.assertEqual(tags, [4, 4, 2])  # Each figure is an image plus a caption box.
        ids = [element.get("id") for element in document.element.body.iter(qn("wp:docPr"))]
        self.assertEqual(len(ids), len(set(ids)))

    def test_after_page_places_figures_between_existing_pages(self):
        source = self.document_with("Page one", "Page two", "Page three", page_breaks=True)
        self.assertEqual(inspect_document(source)["pages_estimate"], 3)
        output = self.folder / "out.docx"
        build_document(source, self.figures[:1], output, insert_mode="after_page", after_page=1)
        document = Document(output)
        body = list(document.element.body)
        figure_page = page_paragraphs(document)[0]._p
        texts = {paragraph.text: body.index(paragraph._p) for paragraph in document.paragraphs
                 if paragraph.text}
        self.assertLess(texts["Page one"], body.index(figure_page))
        self.assertLess(body.index(figure_page), texts["Page two"])
        self.assertEqual(inspect_document(output)["pages_estimate"], 4)

    def test_after_page_zero_puts_figures_first(self):
        output = self.folder / "out.docx"
        build_document(self.document_with("Cover"), self.figures[:1], output,
                       insert_mode="after_page", after_page=0)
        document = Document(output)
        self.assertTrue(document.paragraphs[0]._p.findall(".//" + qn("wp:anchor")))
        self.assertFalse(document.paragraphs[0].paragraph_format.page_break_before)
        self.assertTrue(document.paragraphs[1].paragraph_format.page_break_before)
        self.assertEqual(body_text(document), ["Cover"])

    def test_blank_document_does_not_get_an_empty_first_page(self):
        output = self.folder / "out.docx"
        build_document(self.document_with(""), self.figures[:1], output)
        self.assertEqual(inspect_document(output)["pages_estimate"], 1)

    def test_rendered_break_after_hard_break_counts_once(self):
        document = Document()
        document.add_paragraph("First page")
        document.element.body.insert(1, parse_xml(
            f'<w:p xmlns:w="{NAMESPACES["w"]}"><w:r><w:br w:type="page"/></w:r>'
            f'<w:r><w:lastRenderedPageBreak/><w:t>Second page</w:t></w:r></w:p>'
        ))
        self.assertEqual(inspect_document(io.BytesIO(docx_bytes(document)))["pages_estimate"], 2)

    def test_caption_template_tokens(self):
        figure = Figure(Path("x.png"), "index", ".html", 2, 3)
        self.assertEqual(caption_for("Fig {file_no}.{part}: {file}", figure, 7), "Fig 2.3: index.html")
        with self.assertRaises(ValueError):
            caption_for("{n.__class__}", figure, 1)
        with self.assertRaises(ValueError):
            caption_for("{unknown}", figure, 1)
        with self.assertRaisesRegex(ValueError, "type it twice"):
            caption_for("Fig {n} {", figure, 1)
        with self.assertRaisesRegex(ValueError, "Caption template is not valid"):
            caption_for("{stem:{x}}", figure, 1)
        self.assertEqual(caption_for("{{Fig}} {n}", figure, 4), "{Fig} 4")


class SmartPageTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        patcher = patch.object(web_app, "SMART_ROOT", Path(self.folder.name))
        patcher.start()
        self.addCleanup(patcher.stop)
        self.client = web_app.app.test_client()
        document = Document()
        document.add_paragraph("Report")
        self.docx = docx_bytes(document)

    def build(self, **fields):
        data = {"files": (io.BytesIO(b"<p>hello</p>\n" * 21), "index.html"), "lines": "20"}
        data.update(fields)
        return self.client.post("/api/smart/build", data=data, content_type="multipart/form-data")

    def test_pages_link_to_each_other(self):
        self.assertIn(b'href="/smart"', self.client.get("/").data)
        page = self.client.get("/smart")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"Build Word document", page.data)

    def test_build_then_keep_adding_to_the_same_document(self):
        response = self.build(docx=(io.BytesIO(self.docx), "report.docx"))
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual([item["caption"] for item in response.json["figures"]],
                         ["Fig 1: index1.html", "Fig 2: index2.html"])
        download = self.client.get(response.json["download_url"])
        self.assertEqual(download.status_code, 200)
        self.assertIn("report.docx", download.headers["Content-Disposition"])
        download.close()
        preview = self.client.get(response.json["figures"][0]["image_url"])
        self.assertEqual(preview.status_code, 200)
        preview.close()

        more = self.build(doc_id=response.json["doc_id"], insert_mode="continue",
                          files=(io.BytesIO(b"a { color: red; }\n"), "style.css"))
        self.assertEqual(more.status_code, 200, more.json)
        self.assertEqual(more.json["figures"][0]["caption"], "Fig 3: style1.css")
        self.assertEqual(more.json["total_figures"], 3)

        inspected = self.client.post("/api/smart/inspect", data={
            "doc_id": more.json["doc_id"]}, content_type="multipart/form-data")
        self.assertEqual(inspected.json["next_figure"], 4)

    def test_rejects_documents_that_are_not_docx(self):
        response = self.build(docx=(io.BytesIO(b"hello"), "notes.txt"))
        self.assertEqual(response.status_code, 400)
        broken = self.build(docx=(io.BytesIO(b"not a zip"), "broken.docx"))
        self.assertEqual(broken.status_code, 400)
        self.assertEqual(broken.json["error"], "Could not read this Word document.")
        self.assertEqual(list(Path(self.folder.name).iterdir()), [])

    def test_non_english_names_keep_extensions(self):
        response = self.build(docx=(io.BytesIO(self.docx), "报告.docx"),
                              files=(io.BytesIO(b"<p>hi</p>\n"), "索引.html"))
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json["figures"][0]["caption"], "Fig 1: 索引1.html")
        download = self.client.get(response.json["download_url"])
        self.assertIn("filename*=UTF-8''%E6%8A%A5%E5%91%8A.docx",
                      download.headers["Content-Disposition"])
        download.close()

    def test_bad_caption_template_is_reported_as_a_caption_problem(self):
        response = self.build(docx=(io.BytesIO(self.docx), "report.docx"),
                              caption_template="{stem:{x}}")
        self.assertEqual(response.status_code, 400)
        self.assertIn("Caption template", response.json["error"])

    def test_rejects_unknown_document_ids(self):
        response = self.build(doc_id="0123456789ab")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.client.get("/api/smart/files/../etc/passwd").status_code, 404)

    def test_blank_document_endpoint(self):
        response = self.client.post("/api/smart/blank", data={}, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json["name"], "Document.docx")
        self.assertEqual(response.json["next_figure"], 1)

    def test_build_from_new_document_flag(self):
        response = self.build(new_document="1")
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json["name"], "Document.docx")
        self.assertEqual(response.json["figures"][0]["caption"], "Fig 1: index1.html")


if __name__ == "__main__":
    unittest.main()
